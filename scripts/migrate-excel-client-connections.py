"""
Migra conexiones desde Conexiones_revisado_contiene.xlsx a client_connections.

Solo importa filas con Coincide TeamSupport = SI.

Uso:
  python migrate-excel-client-connections.py              # dry-run (informe)
  python migrate-excel-client-connections.py --apply    # inserta en conexiones_db
"""

from __future__ import annotations

import argparse
import json
import os
import sys
import unicodedata
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import psycopg2

APPSETTINGS = Path(r"c:\Proyectos\Conexiones\src\Conexiones.Web\appsettings.Development.json")
DEFAULT_EXCEL = Path(r"c:\Users\ciro.ascanio\Downloads\Conexiones_revisado_contiene.xlsx")
DEFAULT_REGISTERED_BY = "Migración Excel"

TYPE_ALIASES = {
    "INFORMATICO": "INFORMÁTICO",
    "INFORMÁTICO": "INFORMÁTICO",
}


@dataclass
class MigrationRow:
    excel_row: int
    excel_cliente: str
    customer_id: str
    company_name: str
    razon_social: str | None
    erp_id: str | None
    connection_type_id: int
    connection_type_name: str
    notes: str | None


@dataclass
class MigrationPlan:
    to_insert: list[MigrationRow] = field(default_factory=list)
    skipped: list[tuple[int, str, str]] = field(default_factory=list)


def load_conn_string(key: str) -> dict:
    env_key = "CONEXIONES_CONNECTION" if key == "DefaultConnection" else "TEAMSUPPORT_CONNECTION"
    env = os.environ.get(env_key)
    if env:
        conn_str = env
    else:
        if not APPSETTINGS.exists():
            raise SystemExit(
                f"No se encontró {APPSETTINGS} ni variable {env_key}."
            )
        data = json.loads(APPSETTINGS.read_text(encoding="utf-8"))
        conn_str = data["ConnectionStrings"][key]

    parts: dict[str, str] = {}
    for segment in conn_str.split(";"):
        segment = segment.strip()
        if not segment or "=" not in segment:
            continue
        k, value = segment.split("=", 1)
        parts[k.strip().lower()] = value.strip()

    ssl_required = parts.get("ssl mode", "").lower() == "require"
    return {
        "host": parts.get("host"),
        "port": int(parts.get("port", "5432")),
        "dbname": parts.get("database"),
        "user": parts.get("username"),
        "password": parts.get("password"),
        "sslmode": "require" if ssl_required else "prefer",
    }


def normalize_key(value: str) -> str:
    text = " ".join(value.strip().upper().split())
    text = unicodedata.normalize("NFD", text)
    return "".join(char for char in text if unicodedata.category(char) != "Mn")


def resolve_connection_type(
    excel_value: str,
    type_index: dict[str, tuple[int, str, bool]],
) -> tuple[int, str, bool] | None:
    candidates = [
        normalize_key(excel_value),
        normalize_key(TYPE_ALIASES.get(excel_value.strip().upper(), excel_value)),
    ]
    for key in candidates:
        if key in type_index:
            return type_index[key]
    return None


def pick_customer(
    erp: str,
    company: str,
    by_erp: dict[str, list[tuple[str, str | None, str | None, str | None]]],
    by_company: dict[str, list[tuple[str, str | None, str | None, str | None]]],
) -> tuple[tuple[str, str | None, str | None, str | None] | None, str | None]:
    company_key = normalize_key(company) if company else ""

    if erp:
        matches = by_erp.get(erp, [])
        if len(matches) == 1:
            return matches[0], None
        if len(matches) > 1 and company_key:
            filtered = [
                item
                for item in matches
                if normalize_key(item[1] or "") == company_key
                or normalize_key(item[2] or "") == company_key
            ]
            if len(filtered) == 1:
                return filtered[0], None
            if len(filtered) > 1:
                return None, "cliente_ambiguo_por_erp"
        if len(matches) > 1:
            return None, "cliente_ambiguo_por_erp"

    if company_key:
        matches = by_company.get(company_key, [])
        if len(matches) == 1:
            return matches[0], None
        if len(matches) > 1:
            return None, "cliente_ambiguo_por_nombre"

    return None, "cliente_no_resuelto"


def build_plan(excel_path: Path) -> MigrationPlan:
    if not excel_path.exists():
        raise SystemExit(f"No existe el archivo: {excel_path}")

    conexiones = psycopg2.connect(**load_conn_string("DefaultConnection"))
    teamsupport = psycopg2.connect(**load_conn_string("TeamSupportConnection"))

    cur = conexiones.cursor()
    cur.execute('SELECT "Id", "Name", "IsActive" FROM connection_types ORDER BY "Id"')
    connection_types = cur.fetchall()
    cur.execute('SELECT "TeamSupportClientId", "ConnectionTypeId" FROM client_connections')
    existing_pairs = set(cur.fetchall())
    cur.close()

    ts_cur = teamsupport.cursor()
    ts_cur.execute(
        'SELECT "CustomerId", "CompanyName", "RazonSocial", "ErpId" FROM customers'
    )
    customers = ts_cur.fetchall()
    ts_cur.close()
    teamsupport.close()
    conexiones.close()

    type_index = {normalize_key(name): (type_id, name, is_active) for type_id, name, is_active in connection_types}

    by_erp: dict[str, list[tuple[str, str | None, str | None, str | None]]] = {}
    by_company: dict[str, list[tuple[str, str | None, str | None, str | None]]] = {}
    for customer_id, company, razon, erp in customers:
        record = (str(customer_id), company, razon, str(erp).strip() if erp is not None else None)
        if record[3]:
            by_erp.setdefault(record[3], []).append(record)
        for name in (company, razon):
            if name:
                by_company.setdefault(normalize_key(name), []).append(record)

    wb = openpyxl.load_workbook(excel_path, read_only=True)
    ws = wb.active
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {header: index for index, header in enumerate(headers) if header}
    con_idx = next(i for i, header in enumerate(headers) if header and "conex" in str(header).lower())

    required = ["Cliente", "Coincide TeamSupport", "CompanyName", "RazonSocial", "ErpId DB", "Notes"]
    missing = [name for name in required if name not in idx]
    if missing:
        wb.close()
        raise SystemExit(f"Faltan columnas en el Excel: {', '.join(missing)}")

    plan = MigrationPlan()
    seen_pairs: set[tuple[str, int]] = set()

    for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        coincide = (row[idx["Coincide TeamSupport"]] or "").strip().upper()
        if coincide != "SI":
            plan.skipped.append((row_number, str(row[idx["Cliente"]] or ""), "sin_coincidencia_teamsupport"))
            continue

        excel_cliente = str(row[idx["Cliente"]] or "").strip()
        conexion = str(row[con_idx] or "").strip()
        notes_raw = row[idx["Notes"]]
        notes = str(notes_raw).strip() if notes_raw is not None else None
        company = str(row[idx["CompanyName"]] or "").strip()
        razon = str(row[idx["RazonSocial"]] or "").strip() or None
        erp_raw = row[idx["ErpId DB"]]
        erp = str(erp_raw).strip() if erp_raw is not None else ""

        if not conexion:
            plan.skipped.append((row_number, excel_cliente, "sin_tipo_conexion"))
            continue

        type_info = resolve_connection_type(conexion, type_index)
        if not type_info:
            plan.skipped.append((row_number, excel_cliente, f"tipo_conexion_no_mapeado:{conexion}"))
            continue

        type_id, type_name, is_active = type_info
        if not is_active:
            plan.skipped.append((row_number, excel_cliente, f"tipo_conexion_inactivo:{type_name}"))
            continue

        customer, reason = pick_customer(erp, company, by_erp, by_company)
        if customer is None:
            plan.skipped.append((row_number, excel_cliente, reason or "cliente_no_resuelto"))
            continue

        customer_id, db_company, db_razon, db_erp = customer
        pair = (customer_id, type_id)
        if pair in existing_pairs:
            plan.skipped.append((row_number, excel_cliente, "ya_existe_en_db"))
            continue
        if pair in seen_pairs:
            plan.skipped.append((row_number, excel_cliente, "duplicado_en_excel"))
            continue

        seen_pairs.add(pair)
        plan.to_insert.append(
            MigrationRow(
                excel_row=row_number,
                excel_cliente=excel_cliente,
                customer_id=customer_id,
                company_name=db_company or company,
                razon_social=db_razon or razon,
                erp_id=db_erp or (erp or None),
                connection_type_id=type_id,
                connection_type_name=type_name,
                notes=notes or None,
            )
        )

    wb.close()
    return plan


def print_report(plan: MigrationPlan, excel_path: Path, apply: bool) -> None:
    skipped_stats = Counter(reason.split(":", 1)[0] for _, _, reason in plan.skipped)
    mode = "APLICAR" if apply else "DRY-RUN"

    print(f"Modo: {mode}")
    print(f"Excel: {excel_path}")
    print(f"Insertables: {len(plan.to_insert)}")
    print(f"Omitidas: {len(plan.skipped)}")
    print()
    print("Omitidas por motivo:")
    for reason, count in sorted(skipped_stats.items()):
        print(f"  - {reason}: {count}")

    if plan.to_insert:
        print()
        print("Primeras inserciones previstas:")
        for item in plan.to_insert[:10]:
            print(
                f"  fila {item.excel_row}: {item.excel_cliente} -> "
                f"{item.company_name} [{item.connection_type_name}] "
                f"CustomerId={item.customer_id}"
            )
        if len(plan.to_insert) > 10:
            print(f"  ... y {len(plan.to_insert) - 10} más")

    detailed = [item for item in plan.skipped if "no_mapeado" in item[2] or item[2].startswith("cliente_")]
    if detailed:
        print()
        print("Detalle relevante (omitidas):")
        for row_number, cliente, reason in detailed[:15]:
            print(f"  fila {row_number}: {cliente} [{reason}]")


def apply_plan(plan: MigrationPlan, registered_by: str) -> int:
    if not plan.to_insert:
        print("No hay registros para insertar.")
        return 0

    conn = psycopg2.connect(**load_conn_string("DefaultConnection"))
    inserted = 0
    now = datetime.now(timezone.utc)

    try:
        with conn:
            with conn.cursor() as cur:
                for item in plan.to_insert:
                    cur.execute(
                        """
                        INSERT INTO client_connections (
                            "TeamSupportClientId",
                            "TeamSupportClientName",
                            "TeamSupportClientCode",
                            "TeamSupportClientRazonSocial",
                            "ConnectionTypeId",
                            "Notes",
                            "RegisteredAt",
                            "RegisteredByUserId",
                            "RegisteredByUserName"
                        )
                        VALUES (%s, %s, %s, %s, %s, %s, %s, NULL, %s)
                        """,
                        (
                            item.customer_id,
                            item.company_name[:300],
                            (item.erp_id[:100] if item.erp_id else None),
                            (item.razon_social[:300] if item.razon_social else None),
                            item.connection_type_id,
                            (item.notes[:1000] if item.notes else None),
                            now,
                            registered_by[:256],
                        ),
                    )
                    inserted += 1
    finally:
        conn.close()

    print(f"Insertados: {inserted}")
    return inserted


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Migra client_connections desde Conexiones_revisado_contiene.xlsx"
    )
    parser.add_argument(
        "--excel",
        type=Path,
        default=DEFAULT_EXCEL,
        help=f"Ruta al Excel revisado (default: {DEFAULT_EXCEL})",
    )
    parser.add_argument(
        "--apply",
        action="store_true",
        help="Inserta registros en conexiones_db (por defecto solo informe)",
    )
    parser.add_argument(
        "--registered-by",
        default=DEFAULT_REGISTERED_BY,
        help=f"Valor para RegisteredByUserName (default: {DEFAULT_REGISTERED_BY})",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    plan = build_plan(args.excel)
    print_report(plan, args.excel, args.apply)

    if args.apply:
        print()
        apply_plan(plan, args.registered_by)

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
