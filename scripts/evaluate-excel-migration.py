import json
import os
import sys
from collections import Counter
from pathlib import Path

import openpyxl
import psycopg2

APPSETTINGS = Path(r"c:\Proyectos\Conexiones\src\Conexiones.Web\appsettings.Development.json")
EXCEL_CANDIDATES = [
    Path(r"c:\Users\ciro.ascanio\Downloads\Conexiones_revisado_contiene_v2.xlsx"),
    Path(r"c:\Users\ciro.ascanio\Downloads\Conexiones_revisado_contiene.xlsx"),
]


def load_conn_string(key: str) -> dict:
    env_key = "CONEXIONES_CONNECTION" if key == "DefaultConnection" else "TEAMSUPPORT_CONNECTION"
    env = os.environ.get(env_key)
    if env:
        conn_str = env
    else:
        data = json.loads(APPSETTINGS.read_text(encoding="utf-8"))
        conn_str = data["ConnectionStrings"][key]

    parts = {}
    for segment in conn_str.split(";"):
        segment = segment.strip()
        if not segment or "=" not in segment:
            continue
        k, value = segment.split("=", 1)
        parts[k.strip().lower()] = value.strip()

    return {
        "host": parts.get("host"),
        "port": int(parts.get("port", "5432")),
        "dbname": parts.get("database"),
        "user": parts.get("username"),
        "password": parts.get("password"),
        "sslmode": "require" if parts.get("ssl mode", "").lower() == "require" else "prefer",
    }


def normalize_name(value: str) -> str:
    return " ".join(value.strip().upper().split())


def main() -> int:
    excel_path = next((p for p in EXCEL_CANDIDATES if p.exists()), None)
    if not excel_path:
        print("No se encontró el Excel revisado.", file=sys.stderr)
        return 1

    conexiones = psycopg2.connect(**load_conn_string("DefaultConnection"))
    teamsupport = psycopg2.connect(**load_conn_string("TeamSupportConnection"))

    cur = conexiones.cursor()
    cur.execute(
        'SELECT "Id", "Name", "IsActive" FROM connection_types ORDER BY "Id"'
    )
    connection_types = cur.fetchall()
    cur.execute("SELECT COUNT(*) FROM client_connections")
    existing_count = cur.fetchone()[0]
    cur.execute(
        """
        SELECT "TeamSupportClientId", "ConnectionTypeId"
        FROM client_connections
        """
    )
    existing_pairs = set(cur.fetchall())
    cur.close()

    ts_cur = teamsupport.cursor()
    ts_cur.execute(
        'SELECT "CustomerId", "CompanyName", "RazonSocial", "ErpId" FROM customers'
    )
    customers = ts_cur.fetchall()
    ts_cur.close()
    teamsupport.close()
    conexiones.close()

    by_erp: dict[str, list[tuple]] = {}
    by_company: dict[str, list[tuple]] = {}
    for customer_id, company, razon, erp in customers:
        record = (customer_id, company, razon, erp)
        if erp:
            by_erp.setdefault(str(erp).strip(), []).append(record)
        for name in (company, razon):
            if name:
                by_company.setdefault(normalize_name(name), []).append(record)

    db_type_names = {normalize_name(name): (type_id, name, is_active) for type_id, name, is_active in connection_types}

    wb = openpyxl.load_workbook(excel_path, read_only=True)
    ws = wb.active
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()

    idx = {header: index for index, header in enumerate(headers) if header}
    con_idx = next(i for i, header in enumerate(headers) if header and "conex" in str(header).lower())

    stats = Counter()
    unmapped_types = Counter()
    unresolved_customers = Counter()
    duplicate_rows = Counter()
    ready: list[dict] = []

    for row in rows:
        coincide = (row[idx["Coincide TeamSupport"]] or "").strip().upper()
        if coincide != "SI":
            stats["sin_coincidencia_teamsupport"] += 1
            continue

        excel_cliente = (row[idx["Cliente"]] or "").strip()
        conexion = (row[con_idx] or "").strip()
        notes = (row[idx["Notes"]] or "").strip() if row[idx.get("Notes", -1)] is not None else ""
        company = (row[idx["CompanyName"]] or "").strip()
        razon = (row[idx["RazonSocial"]] or "").strip()
        erp = str(row[idx["ErpId DB"]]).strip() if row[idx["ErpId DB"]] is not None else ""

        if not conexion:
            stats["sin_tipo_conexion"] += 1
            continue

        type_info = db_type_names.get(normalize_name(conexion))
        if not type_info:
            stats["tipo_conexion_no_mapeado"] += 1
            unmapped_types[conexion] += 1
            continue

        type_id, type_name, is_active = type_info
        if not is_active:
            stats["tipo_conexion_inactivo"] += 1
            continue

        matches: list[tuple] = []
        if erp and erp in by_erp:
            matches = by_erp[erp]
        elif company:
            matches = by_company.get(normalize_name(company), [])

        if not matches:
            stats["cliente_no_resuelto_en_teamsupport"] += 1
            unresolved_customers[excel_cliente or company] += 1
            continue

        if len(matches) > 1:
            stats["cliente_ambiguo"] += 1

        customer_id, db_company, db_razon, db_erp = matches[0]
        pair = (customer_id, type_id)
        if pair in existing_pairs:
            stats["ya_existe_en_db"] += 1
            continue

        duplicate_rows[(customer_id, type_id)] += 1
        if duplicate_rows[(customer_id, type_id)] > 1:
            stats["duplicado_en_excel"] += 1
            continue

        ready.append(
            {
                "excel_cliente": excel_cliente,
                "customer_id": customer_id,
                "company_name": db_company,
                "razon_social": db_razon,
                "erp_id": db_erp,
                "connection_type": type_name,
                "notes": notes,
            }
        )
        stats["listo_para_migrar"] += 1

    print(f"Excel: {excel_path.name}")
    print(f"Filas: {len(rows)}")
    print(f"Tipos en conexiones_db: {len(connection_types)}")
    for type_id, name, is_active in connection_types:
        print(f"  - [{type_id}] {name} {'(activo)' if is_active else '(inactivo)'}")
    print(f"Registros actuales en client_connections: {existing_count}")
    print()
    print("Resumen migración:")
    for key in sorted(stats):
        print(f"  {key}: {stats[key]}")
    print()
    if unmapped_types:
        print("Tipos de conexión del Excel sin equivalente en conexiones_db:")
        for name, count in unmapped_types.most_common():
            print(f"  - {name}: {count}")
        print()
    if unresolved_customers:
        print("Filas con coincidencia Excel pero sin CustomerId resoluble:")
        for name, count in unresolved_customers.most_common(10):
            print(f"  - {name}: {count}")
        if len(unresolved_customers) > 10:
            print(f"  ... y {len(unresolved_customers) - 10} más")
        print()
    print(f"Insertables estimados: {len(ready)}")
    if ready[:5]:
        print("Ejemplos:")
        for item in ready[:5]:
            print(
                f"  - {item['excel_cliente']} -> {item['company_name']} "
                f"[{item['connection_type']}] CustomerId={item['customer_id']}"
            )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
