import json
import os
import re
import sys
import unicodedata
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
import psycopg2

EXCEL_IN = Path(r"c:\Users\ciro.ascanio\Downloads\Conexiones (1).xlsx")
EXCEL_OUT = Path(r"c:\Users\ciro.ascanio\Downloads\Conexiones_revisado_contiene.xlsx")
APPSETTINGS = Path(r"c:\Proyectos\Conexiones\src\Conexiones.Web\appsettings.Development.json")

LEGAL_SUFFIXES = re.compile(
    r"\b(s\.?l\.?u?\.?|s\.?a\.?|s\.?l\.?|sociedad limitada|instalaciones)\b",
    re.I,
)
PREFIXES = re.compile(
    r"^(elec\.?|elec:|electricidad|electro[\-\.]?|emo|ing\.?|mont\.?|montaje|sum\.?|sum)\s+",
    re.I,
)
STOP_WORDS = {
    "de", "del", "la", "las", "los", "y", "the", "sl", "sa", "slu",
    "mont", "electr", "elec", "instalaciones", "sistemas", "electricidad",
    "cap", "group", "hi", "hijos", "cano", "nova", "gomez", "pablo",
    "cruz", "valles", "saneamientos", "electr", "mont", "instalaciones",
}
SHORT_COMPACT_BLOCKLIST = {
    "co", "sa", "sl", "de", "del", "elec", "mont", "ema", "ich",
    "cap", "sum", "ing", "emo", "fit", "rel", "ria", "san", "ser",
    "egb", "nova", "cano", "gomez", "pablo",
}


def load_team_support_connection_string() -> str:
    env = os.environ.get("TEAMSUPPORT_CONNECTION")
    if env:
        return env

    if APPSETTINGS.exists():
        data = json.loads(APPSETTINGS.read_text(encoding="utf-8"))
        cs = data.get("ConnectionStrings", {}).get("TeamSupportConnection")
        if cs:
            return cs

    raise SystemExit(
        "Define TEAMSUPPORT_CONNECTION o configura TeamSupportConnection en appsettings.Development.json"
    )


def parse_pg_conn_string(conn_str: str) -> dict:
    parts = {}
    for segment in conn_str.split(";"):
        segment = segment.strip()
        if not segment or "=" not in segment:
            continue
        key, value = segment.split("=", 1)
        parts[key.strip().lower()] = value.strip()

    return {
        "host": parts.get("host"),
        "port": int(parts.get("port", "5432")),
        "dbname": parts.get("database"),
        "user": parts.get("username"),
        "password": parts.get("password"),
        "sslmode": "require" if parts.get("ssl mode", "").lower() == "require" else "prefer",
    }


def normalize(text: str) -> str:
    text = text.strip().lower()
    text = unicodedata.normalize("NFD", text)
    text = "".join(c for c in text if unicodedata.category(c) != "Mn")
    text = re.sub(r"[^\w\s]", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def compact(text: str) -> str:
    return re.sub(r"[^a-z0-9]", "", normalize(text))


def tokens(text: str, min_len: int = 3) -> list[str]:
    return [t for t in normalize(text).split() if len(t) >= min_len and t not in STOP_WORDS]


def strip_all_prefixes(text: str) -> str:
    previous = None
    while previous != text:
        previous = text
        text = PREFIXES.sub("", text).strip()
    return text


def variant_in_text(variant: str, text: str) -> bool:
    variant_l = variant.lower()
    text_l = text.lower()
    if len(variant_l) <= 4:
        return re.search(rf"\b{re.escape(variant_l)}\b", text_l) is not None
    return variant_l in text_l


def search_variants(cliente: str) -> list[tuple[str, bool]]:
    variants: list[tuple[str, bool]] = []
    seen: set[str] = set()

    def add(value: str, from_paren: bool = False) -> None:
        value = value.strip()
        if not value or value in seen:
            return
        seen.add(value)
        variants.append((value, from_paren))

    add(cliente)
    for alias in re.findall(r"\(([^)]+)\)", cliente):
        add(alias.strip(), from_paren=True)

    cleaned = LEGAL_SUFFIXES.sub("", cliente)
    cleaned = re.sub(r"\s*,\s*", " ", cleaned).strip()
    add(cleaned)
    add(strip_all_prefixes(cleaned))
    add(strip_all_prefixes(cliente))

    for token in tokens(cliente, min_len=5):
        add(token)

    words = normalize(cliente).split()
    for size in (3, 2):
        for index in range(len(words) - size + 1):
            chunk_words = words[index : index + size]
            chunk = " ".join(chunk_words)
            chunk_compact = compact(chunk)
            significant = [w for w in chunk_words if w not in STOP_WORDS and len(w) >= 4]
            if not significant:
                if not (len(chunk_words) >= 2 and len(chunk_compact) >= 8 and all(w in STOP_WORDS for w in chunk_words)):
                    continue
            if len(chunk_compact) >= 6:
                add(chunk)

    return variants


def find_contains_only(needle: str, customers: list[tuple]) -> tuple | None:
    needle_l = needle.strip().lower()
    if len(needle_l) < 4:
        return None
    matches = []
    for company, razon, erp in customers:
        company_l = (company or "").strip().lower()
        razon_l = (razon or "").strip().lower()
        if needle_l in company_l or needle_l in razon_l:
            matches.append((company, razon, erp, len(company or "")))
    if not matches:
        return None
    company, razon, erp, _ = min(matches, key=lambda item: item[3])
    return company, razon, erp


def typo_corrections(cliente: str) -> list[str]:
    variants: list[str] = []
    seen: set[str] = set()

    def add(value: str) -> None:
        value = value.strip()
        if value and value.lower() != cliente.lower() and value not in seen:
            seen.add(value)
            variants.append(value)

    normalized = normalize(cliente)
    add(normalized.replace("puesrtol", "puertol"))
    add(normalized.replace("sueprat", "superat"))
    add(cliente.replace("MESANPABLO", "ME SAN PABLO").replace("mesanpablo", "me san pablo"))
    if compact(cliente) == "mesanpablo":
        add("SAN PABLO")
        add("ME SAN PABLO")
    return variants


def pick_best(candidates: list[tuple]) -> tuple | None:
    if not candidates:
        return None
    return min(candidates, key=lambda item: (item[4], item[5], item[6], item[0] or ""))


def find_customer_match(
    cliente: str, customers: list[tuple], *, allow_typo: bool = True
) -> tuple | None:
    cliente = cliente.strip()
    if not cliente:
        return None

    candidates: list[tuple] = []
    cliente_compact = compact(cliente)

    def add_candidate(company, razon, erp, reason: str, rank: int, extra: int = 0) -> None:
        candidates.append((company, razon, erp, reason, rank, extra, len(company or "")))

    needle = cliente.lower()
    for company, razon, erp in customers:
        company_l = (company or "").strip().lower()
        razon_l = (razon or "").strip().lower()
        if needle == company_l or needle == razon_l:
            add_candidate(company, razon, erp, "Coincidencia exacta", 0)
        elif needle in company_l or needle in razon_l:
            add_candidate(company, razon, erp, "Coincidencia contiene", 1)

    if len(cliente_compact) >= 5:
        for company, razon, erp in customers:
            for field in (company, razon):
                field_compact = compact(field or "")
                if len(field_compact) >= 5 and cliente_compact in field_compact:
                    add_candidate(
                        company,
                        razon,
                        erp,
                        "Coincidencia compacta (cliente en DB)",
                        2,
                        len(cliente_compact),
                    )

    cliente_words = set(normalize(cliente).split())

    for variant, from_paren in search_variants(cliente):
        if variant == cliente:
            continue
        variant_l = variant.lower()
        min_variant_len = 3 if from_paren else 5
        if len(variant_l) < min_variant_len:
            continue
        if variant_l in STOP_WORDS:
            continue
        variant_rank = 3 if from_paren or variant_l in cliente_words else 4
        for company, razon, erp in customers:
            company_l = (company or "").strip().lower()
            razon_l = (razon or "").strip().lower()
            if variant_in_text(variant, company_l) or variant_in_text(variant, razon_l):
                add_candidate(
                    company,
                    razon,
                    erp,
                    f"Coincidencia variante: {variant}",
                    variant_rank,
                    -len(variant_l),
                )

        variant_compact = compact(variant)
        variant_words = normalize(variant).split()
        generic_variant = not variant_words or all(
            word in STOP_WORDS or len(word) < 4 for word in variant_words
        )
        if len(variant_compact) >= 6 and not generic_variant:
            for company, razon, erp in customers:
                for field in (company, razon):
                    field_compact = compact(field or "")
                    if len(field_compact) >= 5 and variant_compact in field_compact:
                        add_candidate(
                            company,
                            razon,
                            erp,
                            f"Coincidencia compacta variante: {variant}",
                            5,
                        )

    excel_tokens = tokens(cliente, min_len=5)
    if len(excel_tokens) >= 2:
        for company, razon, erp in customers:
            haystack = normalize(f"{company or ''} {razon or ''}")
            if all(token in haystack for token in excel_tokens):
                add_candidate(
                    company,
                    razon,
                    erp,
                    f"Coincidencia por tokens: {', '.join(excel_tokens)}",
                    5,
                    -len(excel_tokens),
                )

    if len(cliente_compact) >= 6:
        for company, razon, erp in customers:
            for field in (company, razon):
                field_compact = compact(field or "")
                if len(field_compact) < 3 or field_compact in SHORT_COMPACT_BLOCKLIST:
                    continue
                if field_compact not in cliente_compact:
                    continue
                coverage = len(field_compact) / len(cliente_compact)
                if len(field_compact) >= 5 or (len(field_compact) >= 3 and coverage >= 0.25):
                    add_candidate(
                        company,
                        razon,
                        erp,
                        f"Coincidencia compacta inversa: {field}",
                        6,
                        -len(field_compact),
                    )

    if allow_typo:
        typo_variants = typo_corrections(cliente)
        for typo in typo_variants:
            typo_hit = find_customer_match(typo, customers, allow_typo=False)
            if typo_hit:
                add_candidate(
                    typo_hit[0],
                    typo_hit[1],
                    typo_hit[2],
                    f"Coincidencia posible typo: {typo}",
                    8,
                )

    if len(cliente_compact) >= 6:
        for company, razon, erp in customers:
            for field in (company, razon):
                field_compact = compact(field or "")
                if len(field_compact) < 6:
                    continue
                if field_compact[:4] != cliente_compact[:4]:
                    continue
                ratio = SequenceMatcher(None, cliente_compact, field_compact).ratio()
                if ratio >= 0.9:
                    add_candidate(
                        company,
                        razon,
                        erp,
                        f"Coincidencia aproximada ({ratio:.0%})",
                        7,
                        -int(ratio * 100),
                    )

    best = pick_best(candidates)
    if not best:
        return None
    return best[0], best[1], best[2], best[3]


def main() -> int:
    if not EXCEL_IN.exists():
        print(f"No existe el archivo: {EXCEL_IN}", file=sys.stderr)
        return 1

    pg = parse_pg_conn_string(load_team_support_connection_string())

    conn = psycopg2.connect(**pg)
    cur = conn.cursor()
    cur.execute('SELECT "CompanyName", "RazonSocial", "ErpId" FROM customers')
    customers = cur.fetchall()
    cur.close()
    conn.close()

    wb = openpyxl.load_workbook(EXCEL_IN)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    review_cols = [
        "Coincide TeamSupport",
        "CompanyName",
        "RazonSocial",
        "ErpId DB",
        "Motivo",
    ]
    start_col = ws.max_column + 1
    for i, h in enumerate(review_cols):
        cell = ws.cell(row=1, column=start_col + i, value=h)
        cell.font = Font(bold=True)

    green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    for c in range(start_col, start_col + len(review_cols)):
        ws.cell(row=1, column=c).fill = header_fill

    cliente_col = None
    for idx, h in enumerate(headers, start=1):
        if h and str(h).strip().lower() == "cliente":
            cliente_col = idx
            break
    if not cliente_col:
        print("No se encontró columna Cliente", file=sys.stderr)
        return 1

    matched = 0
    unmatched = 0
    empty = 0
    unmatched_names: list[str] = []
    refined_matches: list[tuple[str, str, str]] = []

    for row in range(2, ws.max_row + 1):
        raw = ws.cell(row=row, column=cliente_col).value
        cliente = str(raw).strip() if raw is not None else ""

        if not cliente:
            empty += 1
            values = ["", "", "", "", "Fila sin cliente"]
            fill = None
        else:
            info = find_customer_match(cliente, customers)
            if info:
                matched += 1
                company, razon, erp, reason = info
                if not reason.startswith("Coincidencia contiene") and not reason.startswith(
                    "Coincidencia exacta"
                ):
                    refined_matches.append((cliente, company or "", reason))
                values = ["SI", company or "", razon or "", erp or "", reason]
                fill = green
            else:
                unmatched += 1
                unmatched_names.append(cliente)
                values = [
                    "NO",
                    "",
                    "",
                    "",
                    "Sin coincidencia en CompanyName/RazonSocial",
                ]
                fill = red

        for i, val in enumerate(values):
            cell = ws.cell(row=row, column=start_col + i, value=val)
            if fill and i == 0:
                cell.fill = fill

    ws.auto_filter.ref = ws.dimensions
    try:
        wb.save(EXCEL_OUT)
        output_path = EXCEL_OUT
    except PermissionError:
        fallback = EXCEL_OUT.with_name("Conexiones_revisado_contiene_v2.xlsx")
        wb.save(fallback)
        output_path = fallback

    print(f"Clientes en DB: {len(customers)}")
    print(f"Total filas datos: {ws.max_row - 1}")
    print(f"Coincidencias: {matched}")
    print(f"Sin coincidencia: {unmatched}")
    print(f"Sin cliente: {empty}")
    print(f"Recuperadas con busqueda afinada: {len(refined_matches)}")
    print(f"Guardado: {output_path}")
    if refined_matches:
        print("\nRecuperadas por busqueda afinada:")
        for cliente, company, reason in refined_matches:
            print(f"  - {cliente} -> {company} [{reason}]")
    if unmatched_names:
        print("\nSin coincidencia:")
        for name in unmatched_names:
            print(f"  - {name}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
