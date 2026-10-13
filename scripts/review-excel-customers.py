import json
import os
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
import psycopg2

EXCEL_IN = Path(r"c:\Users\ciro.ascanio\Downloads\Conexiones (1).xlsx")
EXCEL_OUT = Path(r"c:\Users\ciro.ascanio\Downloads\Conexiones_revisado.xlsx")
APPSETTINGS = Path(r"c:\Proyectos\Conexiones\src\Soporte.Web\appsettings.Development.json")


def load_team_support_connection_string() -> str:
    env = os.environ.get("TEAMSUPPORT_CONNECTION")
    if env:
        return env

    if APPSETTINGS.exists():
        data = json.loads(APPSETTINGS.read_text(encoding="utf-8"))
        cs = data.get("ConnectionStrings", {}).get("TeamSupportConnection")
        if cs:
            return cs

    raise SystemExit(
        "Define TEAMSUPPORT_CONNECTION o configura TeamSupportConnection en appsettings.Development.json"
    )


def parse_pg_conn_string(conn_str: str) -> dict:
    parts = {}
    for segment in conn_str.split(";"):
        segment = segment.strip()
        if not segment or "=" not in segment:
            continue
        key, value = segment.split("=", 1)
        parts[key.strip().lower()] = value.strip()

    return {
        "host": parts.get("host"),
        "port": int(parts.get("port", "5432")),
        "dbname": parts.get("database"),
        "user": parts.get("username"),
        "password": parts.get("password"),
        "sslmode": "require" if parts.get("ssl mode", "").lower() == "require" else "prefer",
    }


def main() -> int:
    if not EXCEL_IN.exists():
        print(f"No existe el archivo: {EXCEL_IN}", file=sys.stderr)
        return 1

    pg = parse_pg_conn_string(load_team_support_connection_string())

    conn = psycopg2.connect(**pg)
    cur = conn.cursor()
    cur.execute('SELECT "CompanyName", "RazonSocial", "ErpId" FROM customers')
    rows = cur.fetchall()
    cur.close()
    conn.close()

    name_to_info: dict[str, tuple] = {}

    def add_key(key, company, razon, erp):
        if not key:
            return
        k = key.strip()
        if not k or k in name_to_info:
            return
        name_to_info[k] = (company, razon, erp)

    for company, razon, erp in rows:
        add_key(company, company, razon, erp)
        add_key(razon, company, razon, erp)

    wb = openpyxl.load_workbook(EXCEL_IN)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    review_cols = [
        "Coincide TeamSupport",
        "CompanyName DB",
        "RazonSocial DB",
        "ErpId DB",
        "Motivo",
    ]
    start_col = ws.max_column + 1
    for i, h in enumerate(review_cols):
        cell = ws.cell(row=1, column=start_col + i, value=h)
        cell.font = Font(bold=True)

    green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    for c in range(start_col, start_col + len(review_cols)):
        ws.cell(row=1, column=c).fill = header_fill

    cliente_col = None
    for idx, h in enumerate(headers, start=1):
        if h and str(h).strip().lower() == "cliente":
            cliente_col = idx
            break
    if not cliente_col:
        print("No se encontró columna Cliente", file=sys.stderr)
        return 1

    matched = 0
    unmatched = 0
    empty = 0
    unmatched_names: list[str] = []

    for row in range(2, ws.max_row + 1):
        raw = ws.cell(row=row, column=cliente_col).value
        cliente = str(raw).strip() if raw is not None else ""

        if not cliente:
            empty += 1
            values = ["", "", "", "", "Fila sin cliente"]
            fill = None
        elif cliente in name_to_info:
            matched += 1
            info = name_to_info[cliente]
            values = ["SI", info[0] or "", info[1] or "", info[2] or "", "Coincidencia exacta"]
            fill = green
        else:
            unmatched += 1
            unmatched_names.append(cliente)
            values = ["NO", "", "", "", "Sin coincidencia exacta en CompanyName/RazonSocial"]
            fill = red

        for i, val in enumerate(values):
            cell = ws.cell(row=row, column=start_col + i, value=val)
            if fill and i == 0:
                cell.fill = fill

    ws.auto_filter.ref = ws.dimensions
    wb.save(EXCEL_OUT)

    print(f"Clientes indexados en DB: {len(name_to_info)}")
    print(f"Total filas datos: {ws.max_row - 1}")
    print(f"Coincidencias: {matched}")
    print(f"Sin coincidencia: {unmatched}")
    print(f"Sin cliente: {empty}")
    print(f"Guardado: {EXCEL_OUT}")
    if unmatched_names:
        print("\nPrimeros 25 sin coincidencia:")
        for name in unmatched_names[:25]:
            print(f"  - {name}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
